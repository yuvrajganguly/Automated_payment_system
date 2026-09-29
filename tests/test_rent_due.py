"""RENT_DUE — the week's rent for riders no payout will ever cover.

Until September 2026 a Zomato or Elastic holder's week did not exist in
``transactions``: no file, no cycle, nothing booked, while Raft's cost for the
same days landed in the day-ledger. These tests pin the rule that fixed it —
a rider owes rent for holding the vehicle, whoever pays him — and the seams
around it: only direct-pay people, one week at a time, never twice, recovered
by the ordinary cash path, and a per-id override for a client that pays two
ways at once.
"""

from __future__ import annotations

from datetime import date

import pytest

from payout.domain.payment_model import person_is_direct_pay
from payout.domain.rent_due import book_rent_due, last_week_end, scan_rent_due, week_bounds
from tests.conftest import assign, make_ev, make_person, make_rider

SUNDAY = "2026-09-20"  # W38 as Raft bills it: Mon 14 .. Sun 20
TODAY = date(2026, 9, 21)


def _holder(db, name, rider_id, company, ev, *, handover="2026-09-01", model="Blue"):
    pid = make_person(db, name)
    make_rider(db, pid, rider_id, company, name)
    assign(db, pid, make_ev(db, ev, provider="Raft", model=model), handover=handover)
    return pid


def _rows(db, pid, et="RENT_DUE"):
    return db.execute(
        "SELECT amount, days, cycle_start, cycle_end, remarks FROM transactions "
        "WHERE person_id=? AND event_type=? ORDER BY id",
        (pid, et),
    ).fetchall()


def test_a_zomato_holder_owes_the_week_and_a_jiffy_holder_does_not(db):
    zom = _holder(db, "Direct Das", "Z-1", "Zomato", "CBICEVD9001")
    jif = _holder(db, "File Fatima", "J-1", "Jiffy", "CBICEVD9002")
    db.commit()
    assert person_is_direct_pay(db, zom) and not person_is_direct_pay(db, jif)
    found = scan_rent_due(db, SUNDAY, today=TODAY)
    assert [e["person_id"] for e in found] == [zom]
    e = found[0]
    assert (e["day_from"], e["day_to"], e["days"]) == ("2026-09-14", "2026-09-20", 7)
    assert e["amount"] == 129_500  # a full Blue week at our rate
    assert not e["already_booked"]


def test_booking_writes_rent_due_to_arrears_and_marks_the_days(db):
    pid = _holder(db, "Direct Das", "Z-1", "Zomato", "CBICEVD9001")
    db.commit()
    booked = book_rent_due(db, SUNDAY, created_by="t@t", today=TODAY)
    assert len(booked) == 1 and booked[0]["amount"] == 129_500
    (row,) = _rows(db, pid)
    assert row["amount"] == -129_500 and row["days"] == 7
    assert (row["cycle_start"], row["cycle_end"]) == ("2026-09-14", "2026-09-20")
    assert "collect in cash" in row["remarks"]
    assert _rows(db, pid, "RENT_MISSED") == []  # present, not absent
    arr = db.execute("SELECT outstanding FROM ev_arrears WHERE person_id=?", (pid,)).fetchone()
    assert arr["outstanding"] == 129_500
    days = db.execute(
        "SELECT day, billing_status, daily_cost, provider_cost FROM ev_daily_ledger "
        "WHERE ev_id='CBICEVD9001' AND day BETWEEN '2026-09-14' AND '2026-09-20' ORDER BY day"
    ).fetchall()
    assert len(days) == 7 and {d["billing_status"] for d in days} == {"due"}
    assert days[0]["daily_cost"] == 18_500 and days[0]["provider_cost"] == 17_500  # ours vs Raft's
    meter = db.execute(
        "SELECT rent_charged_through FROM ev_assignments WHERE person_id=?", (pid,)
    ).fetchone()
    assert str(meter["rent_charged_through"])[:10] == "2026-09-20"


def test_booking_twice_books_nothing_twice(db):
    pid = _holder(db, "Direct Das", "Z-1", "Zomato", "CBICEVD9001")
    db.commit()
    assert len(book_rent_due(db, SUNDAY, created_by="t@t", today=TODAY)) == 1
    assert book_rent_due(db, SUNDAY, created_by="t@t", today=TODAY) == []
    assert len(_rows(db, pid)) == 1
    again = scan_rent_due(db, SUNDAY, today=TODAY)
    assert again == [] or all(e["already_booked"] for e in again)


def test_the_next_week_starts_where_the_meter_stopped(db):
    """One week at a time: the historic gap before the first booking is not
    reached back for (the office waived those on 2026-09-04), and the week
    after picks up from the meter."""
    pid = _holder(db, "Direct Das", "Z-1", "Zomato", "CBICEVD9001", handover="2026-08-01")
    db.commit()
    book_rent_due(db, SUNDAY, created_by="t@t", today=TODAY)
    nxt = scan_rent_due(db, "2026-09-27", today=date(2026, 9, 28))
    assert [(e["day_from"], e["day_to"]) for e in nxt] == [("2026-09-21", "2026-09-27")]
    assert _rows(db, pid)[0]["cycle_start"] == "2026-09-14"  # nothing before it


def test_a_handover_mid_week_owes_only_the_days_after_it(db):
    _holder(db, "Direct Das", "Z-1", "Zomato", "CBICEVD9001", handover="2026-09-15")
    db.commit()
    (e,) = scan_rent_due(db, SUNDAY, today=TODAY)
    assert (e["day_from"], e["day_to"], e["days"]) == ("2026-09-16", "2026-09-20", 5)
    assert e["amount"] == 92_500


def test_a_rider_id_can_override_the_company(db):
    """Zomato programme 2 pays through us: mark that id payout_file and the
    person stops being direct-pay, so his rent waits for the cycle."""
    pid = _holder(db, "Two Ways", "Z-2", "Zomato", "CBICEVD9001")
    db.commit()
    assert person_is_direct_pay(db, pid)
    db.execute("UPDATE rider_master SET payment_model='payout_file' WHERE rider_id='Z-2'")
    db.commit()
    assert not person_is_direct_pay(db, pid)
    assert scan_rent_due(db, SUNDAY, today=TODAY) == []
    # and the other way round: a Jiffy id the client pays directly
    db.execute(
        "UPDATE rider_master SET payment_model='direct', company='Jiffy' WHERE rider_id='Z-2'"
    )
    db.commit()
    assert person_is_direct_pay(db, pid)


def test_a_second_id_at_a_file_company_means_the_cycle_will_bill_him(db):
    pid = _holder(db, "Both", "Z-1", "Zomato", "CBICEVD9001")
    make_rider(db, pid, "J-9", "Jiffy", "Both")
    db.commit()
    assert not person_is_direct_pay(db, pid)
    assert scan_rent_due(db, SUNDAY, today=TODAY) == []


def test_week_helpers():
    assert week_bounds("2026-09-20") == (date(2026, 9, 14), date(2026, 9, 20))
    assert last_week_end(date(2026, 9, 21)) == date(2026, 9, 20)  # Monday → yesterday
    assert last_week_end(date(2026, 9, 27)) == date(2026, 9, 20)  # Sunday → last Sunday
    assert last_week_end(date(2026, 9, 26)) == date(2026, 9, 20)


@pytest.fixture
def client(db):
    from fastapi.testclient import TestClient

    from payout.api.app import app
    from payout.db import get_connection

    conn = get_connection()
    try:
        conn.execute(
            "INSERT INTO users (email, password_hash, role, is_active) VALUES (?, ?, 'creator', 1)",
            (
                "boss@t.test",
                __import__("payout.auth", fromlist=["x"]).hash_password("Creator-pass-1"),
            ),
        )
        conn.commit()
    finally:
        conn.close()
    with TestClient(app) as c:
        yield c


def _hdr(client, email="boss@t.test", pw="Creator-pass-1"):
    r = client.post("/api/auth/login", data={"username": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def test_preview_book_cash_and_the_collection_sheet_over_http(db, client):
    pid = _holder(db, "Direct Das", "Z-1", "Zomato", "CBICEVD9001")
    db.execute("UPDATE rider_master SET hub='Ruby', mob_no='9000000001' WHERE rider_id='Z-1'")
    db.commit()
    boss = _hdr(client)

    pv = client.get("/api/rent-due/preview", params={"week_end": SUNDAY}, headers=boss)
    assert pv.status_code == 200, pv.text
    assert pv.json()["pending_count"] == 1 and pv.json()["amount"] == 1295.0  # rupees at the edge
    assert pv.json()["entries"][0]["weekly_rate"] == 1295.0

    bad = client.get("/api/rent-due/preview", params={"week_end": "2026-09-19"}, headers=boss)
    assert bad.status_code == 400  # a Saturday is not a week end

    bk = client.post("/api/rent-due/book", json={"week_end": SUNDAY}, headers=boss)
    assert bk.status_code == 200, bk.text
    assert bk.json()["count"] == 1 and bk.json()["amount"] == 1295.0
    assert (
        client.post("/api/rent-due/book", json={"week_end": SUNDAY}, headers=boss).json()["count"]
        == 0
    )

    # the sheet the hub gets on Monday
    xl = client.post("/api/rent-due/collection", json={"week_end": SUNDAY}, headers=boss)
    assert xl.status_code == 200 and xl.content[:2] == b"PK"
    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(xl.content)).active
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]], strict=True))
    assert row["Name"] == "Direct Das" and row["Hub"] == "Ruby" and row["Phone"] == "9000000001"
    assert (
        row["Rent this week"] == 1295.0 and row["To collect"] == 1295.0 and row["Booked?"] == "yes"
    )

    # the rider pays cash at the hub: the ordinary manual-payment path clears it
    pay = client.post(
        "/api/ledger/rent-payment",
        json={"person_id": pid, "amount": 1295, "paid_on": "2026-09-22"},
        headers=boss,
    )
    assert pay.status_code == 200, pay.text
    arr = db.execute("SELECT outstanding FROM ev_arrears WHERE person_id=?", (pid,)).fetchone()
    assert arr["outstanding"] == 0
    assert _rows(db, pid, "RENT_RECOVERED")  # recovered, arrears first
    # and the week reads as charged AND collected everywhere that sums the ledger
    tot = db.execute(
        "SELECT SUM(CASE WHEN event_type='RENT_DUE' THEN -amount END) AS due, "
        "       SUM(CASE WHEN event_type='RENT_RECOVERED' THEN amount END) AS rec "
        "FROM transactions WHERE person_id=?",
        (pid,),
    ).fetchone()
    assert tot["due"] == 129_500 and tot["rec"] == 129_500


def test_a_recruiter_cannot_see_or_book_rent_due(db, client):
    from payout.auth import hash_password

    db.execute(
        "INSERT INTO users (email, password_hash, role, is_active) VALUES (?, ?, 'recruiter', 1)",
        ("rec@t.test", hash_password("Recruit-pass-1")),
    )
    db.commit()
    h = _hdr(client, "rec@t.test", "Recruit-pass-1")
    assert client.get("/api/rent-due/preview", headers=h).status_code == 403
    assert client.post("/api/rent-due/book", json={}, headers=h).status_code == 403


def test_payment_model_on_the_rider_over_http(db, client):
    pid = _holder(db, "Two Ways", "Z-2", "Zomato", "CBICEVD9001")
    db.commit()
    boss = _hdr(client)
    r = client.get("/api/riders", params={"q": "Two Ways"}, headers=boss)
    assert r.status_code == 200 and r.json()[0]["payment_model"] is None
    p = client.patch(
        "/api/riders/Z-2",
        params={"company": "Zomato"},
        json={"payment_model": "payout_file"},
        headers=boss,
    )
    assert p.status_code == 200, p.text
    assert p.json()["payment_model"] == "payout_file"
    assert not person_is_direct_pay(db, pid)
    bad = client.patch(
        "/api/riders/Z-2",
        params={"company": "Zomato"},
        json={"payment_model": "cheque"},
        headers=boss,
    )
    assert bad.status_code == 400
    cleared = client.patch(
        "/api/riders/Z-2", params={"company": "Zomato"}, json={"payment_model": ""}, headers=boss
    )
    assert cleared.status_code == 200 and cleared.json()["payment_model"] is None
