"""The house sheet: our own four columns for a company that sends no file,
run through the ordinary cycle; and the collection sheet coming back from
the hub with cash receipts filled in.
"""

from __future__ import annotations

import io

import pytest

from tests.conftest import assign, make_ev, make_person, make_rider


@pytest.fixture
def client(db):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    from payout.api import ratelimit
    from payout.api.app import app
    from payout.auth import hash_password

    db.execute(
        "INSERT INTO users (email, password_hash, role, is_active) VALUES (?,?,?,1)",
        ("adm@t.test", hash_password("Admin-pass-1"), "admin"),
    )
    db.commit()
    ratelimit.reset()
    with TestClient(app) as c:
        r = c.post("/api/auth/login", data={"username": "adm@t.test", "password": "Admin-pass-1"})
        assert r.status_code == 200
        yield c


def _sheet(rows, headers=("Person ID", "Name", "Gross payout", "Rent charged", "Net payout")):
    import pandas as pd

    buf = io.BytesIO()
    pd.DataFrame(rows, columns=list(headers)).to_excel(buf, index=False)
    return buf.getvalue()


def _house_company(client, name="Zomato"):
    r = client.patch(
        f"/api/companies/{name}", json={"payment_model": "payout_file", "parser_type": "house"}
    )
    assert r.status_code == 200, r.text
    assert r.json()["parser_type"] == "house"


def _run(client, company, file_bytes, *, commit=False):
    return client.post(
        "/api/cycles/run",
        data={
            "company": company,
            "cycle_start": "2026-09-21",
            "cycle_end": "2026-09-27",
            "commit": "true" if commit else "false",
        },
        files={"file": ("zomato_w39.xlsx", file_bytes)},
    )


def test_the_house_sheet_runs_like_a_client_file(db, client):
    ev_holder = make_person(db, "Holder Halder")
    make_rider(db, ev_holder, "Z-1", "Zomato", "Holder Halder")
    assign(
        db,
        ev_holder,
        make_ev(db, "CBICEVD9001", provider="Raft", model="Blue"),
        handover="2026-09-01",
    )
    walker = make_person(db, "Bike Bose")
    make_rider(db, walker, "Z-2", "Zomato", "Bike Bose")
    db.commit()
    _house_company(client)

    r = _run(
        client,
        "Zomato",
        _sheet(
            [[ev_holder, "Holder Halder", 5000, 1295, 3705], [walker, "Bike Bose", 4200, 0, 4200]]
        ),
    )
    assert r.status_code == 200, r.text
    res = r.json()["result"]
    rows = {x["rider_id"]: x for x in res["pay_rows"] + res["dues_rows"]}
    assert rows["Z-1"]["payout"] == 5000.0 and rows["Z-1"]["rent"] == 1295.0
    assert rows["Z-2"]["payout"] == 4200.0 and rows["Z-2"]["rent"] == 0
    assert rows["Z-1"]["released"] == 3705.0 and rows["Z-2"]["released"] == 4200.0
    assert not any("sheet says" in w for w in res["warnings"])  # sheet agreed with the engine

    # commit: PAYOUT and RENT rows exist for a company that never sent a file
    r = _run(
        client,
        "Zomato",
        _sheet(
            [[ev_holder, "Holder Halder", 5000, 1295, 3705], [walker, "Bike Bose", 4200, 0, 4200]]
        ),
        commit=True,
    )
    assert r.status_code == 200, r.text
    ev = db.execute(
        "SELECT event_type, amount FROM transactions WHERE person_id=? ORDER BY id", (ev_holder,)
    ).fetchall()
    types = {e["event_type"] for e in ev}
    assert {"PAYOUT", "RENT", "RENT_COLLECTED"} <= types
    assert (
        db.execute(
            "SELECT COUNT(*) AS n FROM transactions WHERE person_id=?", (walker,)
        ).fetchone()["n"]
        >= 1
    )


def test_the_sheets_rent_is_checked_not_obeyed(db, client):
    pid = make_person(db, "Holder Halder")
    make_rider(db, pid, "Z-1", "Zomato", "Holder Halder")
    assign(
        db, pid, make_ev(db, "CBICEVD9001", provider="Raft", model="Blue"), handover="2026-09-01"
    )
    db.commit()
    _house_company(client)
    r = _run(client, "Zomato", _sheet([[pid, "Holder Halder", 5000, 1110, 3890]]))
    assert r.status_code == 200, r.text
    res = r.json()["result"]
    row = (res["pay_rows"] + res["dues_rows"])[0]
    assert row["rent"] == 1295.0  # the engine's, from seven held days
    assert any(
        "sheet says rent ₹1,110" in w and "engine charged ₹1,295" in w for w in res["warnings"]
    )
    assert any("sheet says net ₹3,890" in w for w in res["warnings"])


def test_a_person_with_no_id_at_the_company_is_refused_by_name(db, client):
    pid = make_person(db, "Nobody Here")
    make_rider(db, pid, "J-1", "Jiffy", "Nobody Here")  # an id, but at Jiffy
    db.commit()
    _house_company(client)
    r = _run(client, "Zomato", _sheet([[pid, "Nobody Here", 5000, 0, 5000]]))
    assert r.status_code == 400
    assert (
        "no active rider id at Zomato" in r.json()["detail"] and "Nobody Here" in r.json()["detail"]
    )


def test_a_rider_id_column_wins_and_the_total_row_is_ignored(db, client):
    pid = make_person(db, "Two Ids")
    make_rider(db, pid, "Z-old", "Zomato", "Two Ids")
    make_rider(db, pid, "Z-new", "Zomato", "Two Ids")
    db.commit()
    _house_company(client)
    sheet = _sheet(
        [[pid, "Z-new", 5000, 0, 5000], ["", "TOTAL", 5000, 0, 5000]],
        headers=("Person ID", "Rider ID", "Gross payout", "Rent charged", "Net payout"),
    )
    r = _run(client, "Zomato", sheet)
    assert r.status_code == 200, r.text
    res = r.json()["result"]
    assert [x["rider_id"] for x in res["pay_rows"] + res["dues_rows"]] == ["Z-new"]


def test_the_collection_sheet_comes_back_as_receipts(db, client):
    from payout.domain.rent_due import book_rent_due

    a = make_person(db, "Pays Paul")
    make_rider(db, a, "Z-1", "Zomato", "Pays Paul")
    assign(db, a, make_ev(db, "CBICEVD9001", provider="Raft", model="Blue"), handover="2026-09-01")
    b = make_person(db, "Owes Omar")
    make_rider(db, b, "Z-2", "Zomato", "Owes Omar")
    assign(db, b, make_ev(db, "CBICEVD9002", provider="Raft", model="Blue"), handover="2026-09-01")
    db.commit()
    from datetime import date

    book_rent_due(db, "2026-09-20", created_by="t@t", today=date(2026, 9, 21))
    db.commit()

    xl = client.post("/api/rent-due/collection", json={"week_end": "2026-09-20"})
    assert xl.status_code == 200
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xl.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    ci, by = header.index("Collected ₹ (fill in)") + 1, header.index("Collected by") + 1
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Pays Paul":
            ws.cell(row=row[0].row, column=ci, value=1295)
            ws.cell(row=row[0].row, column=by, value="Dipankar")
    out = io.BytesIO()
    wb.save(out)

    up = client.post(
        "/api/rent-due/collection/upload", files={"file": ("collect.xlsx", out.getvalue())}
    )
    assert up.status_code == 200, up.text
    j = up.json()
    assert j["count"] == 1 and j["amount"] == 1295.0 and j["skipped"] == 1 and j["failed"] == []
    assert (
        db.execute("SELECT outstanding FROM ev_arrears WHERE person_id=?", (a,)).fetchone()[
            "outstanding"
        ]
        == 0
    )
    assert (
        db.execute("SELECT outstanding FROM ev_arrears WHERE person_id=?", (b,)).fetchone()[
            "outstanding"
        ]
        == 129_500
    )
    rem = db.execute(
        "SELECT remarks FROM transactions WHERE person_id=? AND event_type='RENT_RECOVERED'", (a,)
    ).fetchone()["remarks"]
    assert "Collection sheet" in rem and "Dipankar" in rem and "CBICEVD9001" in rem
