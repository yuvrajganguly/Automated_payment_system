"""Shadowfax bulk payout: many cumulative Vendor_data files, the pincode
ratecard, the already-paid guard — and the spec's acceptance table.

Fixture: tests/shadowfax_fixture.py, the ten September 2026 downloads with the
riders anonymised (Alpha = the spec's 28307084, Bravo = 2665213, Charlie =
28489235, Delta = 3438374, Echo = 28393260). Every rupee figure below is the
spec's hand-checked one.
"""

from __future__ import annotations

import io
import json
from datetime import date

import pytest

from payout.domain import shadowfax as sf
from tests import shadowfax_fixture as fx
from tests.conftest import make_person, make_rider

ALPHA, BRAVO, CHARLIE, DELTA, ECHO = "91000001", "91000002", "91000003", "91000004", "91000005"
C17, C27 = date(2026, 9, 17), date(2026, 9, 27)


def _price(db, files, cs=C17, ce=C27, fallback=1400):
    return sf.price(db, "Shadowfax", sf.read_files(files), cs, ce, fallback_rate=fallback)


def _by(p):
    return {r["rider_id"]: r for r in p.riders}


# ── the acceptance table ────────────────────────────────────────────────────


def test_the_spec_table_to_the_rupee(db):
    r = _by(_price(db, fx.all_files()))
    expect = {  # days, orders, ppd, cod, rvp, sdd, club, gross ₹, shadowfax ₹
        ALPHA: (4, 90, 12, 56, 8, 0, 14, 1150, 1240),
        BRAVO: (1, 4, 2, 2, 0, 0, 0, 54, 58),
        CHARLIE: (2, 23, 3, 14, 0, 4, 2, 321, 344),
        DELTA: (2, 75, 7, 55, 7, 0, 6, 1047, 1122),
        ECHO: (0, 0, 0, 0, 0, 0, 0, 0, 0),
    }
    for rid, e in expect.items():
        x = r[rid]
        got = (
            x["days"],
            x["orders"],
            x["ppd_orders"],
            x["cod_orders"],
            x["rvp_orders"],
            x["sdd_orders"],
            x["club_orders"],
            x["gross"] // 100,
            x["sfx_payout"] // 100,
        )
        assert got == e, (rid, got, e)
    p = _price(db, fx.all_files())
    assert (p.totals["orders"], p.totals["gross"], p.totals["sfx_payout"]) == (
        192,
        257_200,
        276_400,
    )
    assert ECHO not in {rec.rider_id for rec in p.records}  # zero orders: ignored, not an error


def test_1_dedupe_the_last_file_alone_equals_all_ten(db):
    last = [f for f in fx.all_files() if f[0].endswith("09-28.xlsx")]
    a, b = _price(db, last), _price(db, fx.all_files())
    key = lambda p: sorted((r["rider_id"], r["orders"], r["gross"]) for r in p.riders)  # noqa: E731
    assert key(a) == key(b)


def test_2_a_revised_day_takes_the_newest_file(db):
    def alpha_17(files):
        p = _price(db, files)
        return [
            x["orders"]
            for x in p.detail
            if x["rider_id"] == ALPHA
            and x["order_date"] == "2026-09-17"
            and x["pincode"] == "700033"
        ]

    assert alpha_17([f for f in fx.all_files() if f[0].endswith("09-18.xlsx")]) == [10]
    assert alpha_17(fx.all_files()) == [11]
    # upload order does not matter: the date in the name decides
    assert alpha_17(list(reversed(fx.all_files()))) == [11]


def test_3_pincode_rates(db):
    p = _price(db, fx.all_files())
    rates = lambda pred: {  # noqa: E731
        (x["rate_ppd"], x["rate_cod"], x["rate_rvp"], x["rate_sdd"], x["rate_club"])
        for x in p.detail
        if pred(x)
    }
    assert rates(lambda x: x["pincode"] == "700154") == {(1400, 1500, 1500, 1900, 700)}
    assert rates(lambda x: x["pincode"] != "700154") == {(1300, 1400, 1400, 1800, 700)}
    assert {r["pincode"] for r in p.ratecard_used} == {x["pincode"] for x in p.detail}


def test_4_cycle_filter(db):
    p = _price(db, fx.all_files(), date(2026, 9, 25), C27)
    assert sorted({x["order_date"] for x in p.detail}) == ["2026-09-25", "2026-09-26", "2026-09-27"]


def test_7_missing_pincode_falls_back_and_is_flagged(db):
    rows = [
        (
            799999,
            "P1",
            ALPHA,
            "Rider Alpha",
            "CCU_Dhakuria",
            "2026-09-20",
            50,
            5,
            1,
            2,
            1,
            0,
            1,
            0,
            0,
            0,
        )
    ]
    p = _price(db, [("Vendor_data_2026-09-21.xlsx", fx.xlsx("x", rows))])
    (line,) = p.detail
    assert line["rider_pay"] == 4 * 1400 + 1 * 700 and "rate fallback" in line["flags"]
    assert any("fallback" in w for w in p.warnings)


def test_unrated_orders_are_paid_nothing_and_flagged(db):
    rows = [
        (
            700033,
            "P1",
            ALPHA,
            "Rider Alpha",
            "CCU_Dhakuria",
            "2026-09-20",
            50,
            6,
            1,
            1,
            0,
            0,
            0,
            2,
            2,
            1,
        )
    ]
    p = _price(db, [("Vendor_data_2026-09-21.xlsx", fx.xlsx("x", rows))])
    (line,) = p.detail
    assert line["rider_pay"] == 1300 + 1400
    assert any(f.startswith("unrated") for f in line["flags"])


def test_a_file_without_a_date_in_its_name_is_dated_by_its_content(db):
    b = sf.read_files(
        [
            ("download.xlsx", fx.xlsx("Vendor_data_2026-09-28.xlsx")),
            ("Vendor_data_2026-09-18.xlsx", fx.xlsx("Vendor_data_2026-09-18.xlsx")),
        ]
    )
    assert [f.file_date for f in b.files] == ["2026-09-18", "2026-09-27"]
    assert b.files[1].dated_by == "content"


def test_hub_names_normalise():
    assert sf.norm_hub("CCU_Narendrapur") == sf.norm_hub("Narendra Pur") == "narendrapur"
    assert sf.norm_hub(float("nan")) == ""


# ── over HTTP: the whole cycle ──────────────────────────────────────────────


@pytest.fixture
def client(db):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    from payout.api import ratelimit
    from payout.api.app import app
    from payout.auth import hash_password

    db.execute(
        "INSERT INTO users (email, password_hash, role, is_active) VALUES (?,?,?,1)",
        ("adm@t.test", hash_password("Admin-pass-1"), "admin"),
    )
    db.commit()
    ratelimit.reset()
    with TestClient(app) as c:
        assert (
            c.post(
                "/api/auth/login", data={"username": "adm@t.test", "password": "Admin-pass-1"}
            ).status_code
            == 200
        )
        yield c


def _roster(db, *, bank=True):
    pids = {}
    for rid, name, hub in (
        (ALPHA, "Rider Alpha", "Dhakuria"),
        (BRAVO, "Rider Bravo", "Dhakuria"),
        (CHARLIE, "Rider Charlie", "Dhakuria"),
        (DELTA, "Rider Delta", "Narendra Pur"),
    ):
        pid = make_person(db, name)
        make_rider(db, pid, rid, "Shadowfax", name)
        db.execute(
            "UPDATE rider_master SET hub=?, account_no=? WHERE rider_id=?",
            (hub, "1234567890" if bank else None, rid),
        )
        pids[rid] = pid
    db.commit()
    return pids


def _run(client, files, cs="2026-09-17", ce="2026-09-27", *, commit=False, **extra):
    return client.post(
        "/api/cycles/run",
        data={
            "company": "Shadowfax",
            "cycle_start": cs,
            "cycle_end": ce,
            "commit": "true" if commit else "false",
            **extra,
        },
        files=[("files", (n, b)) for n, b in files],
    )


def test_preview_over_http_matches_the_table(db, client):
    _roster(db)
    r = _run(client, fx.all_files())
    assert r.status_code == 200, r.text
    res, s = r.json()["result"], r.json()["shadowfax"]
    rows = {x["rider_id"]: x for x in res["pay_rows"] + res["dues_rows"]}
    assert {k: rows[k]["payout"] for k in rows} == {
        ALPHA: 1150.0,
        BRAVO: 54.0,
        CHARLIE: 321.0,
        DELTA: 1047.0,
    }
    assert (
        s["totals"]["gross"] == 2572.0
        and s["totals"]["sfx_payout"] == 2764.0
        and s["totals"]["margin"] == 192.0
    )
    assert [f["file_date"] for f in s["files"]] == [
        "2026-09-18",
        "2026-09-19",
        "2026-09-20",
        "2026-09-21",
        "2026-09-22",
        "2026-09-24",
        "2026-09-25",
        "2026-09-26",
        "2026-09-27",
        "2026-09-28",
    ]
    found = {d["order_date"]: d for d in s["order_dates"]}
    assert found["2026-09-17"]["source_file"] == "Vendor_data_2026-09-28.xlsx"
    assert not found["2026-09-18"]["found"]  # missing days are a warning, not an error
    assert any("No Shadowfax data" in w for w in res["warnings"])
    # the file's hub never lands on the roster
    hubs = dict(
        db.execute("SELECT rider_id, hub FROM rider_master WHERE company='Shadowfax'").fetchall()
    )
    assert hubs[DELTA] == "Narendra Pur"


def test_5_already_paid_days_are_left_out(db, client):
    _roster(db)
    r = _run(client, fx.all_files(), "2026-09-17", "2026-09-23", commit=True)
    assert r.status_code == 200, r.text
    assert r.json()["xlsx"]["filename"]
    paid = db.execute("SELECT DISTINCT order_date FROM payout_order_days ORDER BY 1").fetchall()
    assert [p["order_date"] for p in paid] == ["2026-09-17", "2026-09-23"]

    r = _run(client, fx.all_files(), "2026-09-17", "2026-09-27")
    assert r.status_code == 200, r.text
    s = r.json()["shadowfax"]
    assert sorted({a["order_date"] for a in s["already_paid"]}) == ["2026-09-17", "2026-09-23"]
    assert all(a["cycle"] == "2026-09-17..2026-09-23" for a in s["already_paid"])
    got = sum(x["gross"] for x in s["riders"])
    paid_before = sum(a["paid_gross"] for a in s["already_paid"])
    assert got + paid_before == 2572.0  # nothing paid twice, nothing lost


def test_6_an_unknown_rider_is_listed_with_its_pay(db, client):
    _roster(db)
    db.execute("DELETE FROM rider_master WHERE rider_id=?", (DELTA,))
    db.commit()
    r = _run(client, fx.all_files())
    assert r.status_code == 200, r.text
    unk = {u["rider_id"]: u for u in r.json()["result"]["unknown_riders"]}
    assert unk[DELTA]["name"] == "Rider Delta" and unk[DELTA]["gross"] == 1047.0
    assert unk[DELTA]["hub"] == "" or unk[DELTA]["hub"] == "Narendra Pur"


def test_a_rider_with_no_bank_account_is_held(db, client):
    _roster(db, bank=False)
    r = _run(client, fx.all_files())
    res = r.json()["result"]
    rows = res["pay_rows"] + res["dues_rows"]
    assert rows and all(
        x["is_hold"] for x in rows
    )  # the engine's HOLD: marked, not sent to the bank
    assert {x["bank"] for x in r.json()["shadowfax"]["riders"] if x["orders"]} == {"missing — held"}


def test_ztp_is_off_unless_chosen(db, client):
    pids = _roster(db)
    files = fx.all_files()
    files[-1] = (
        files[-1][0],
        fx.xlsx(
            files[-1][0],
            ztp=[
                {
                    "rider_id": ALPHA,
                    "parent_id": 1,
                    "awb_number": "AWB1",
                    "fraud_type": "fake",
                    "total_penalty_amount": 200,
                    "created_date": "2026-09-26",
                },
                {
                    "rider_id": ALPHA,
                    "parent_id": 1,
                    "awb_number": "AWB1",
                    "fraud_type": "fake",
                    "total_penalty_amount": 200,
                    "created_date": "2026-09-26",
                },
            ],
        ),
    )
    r = _run(client, files)
    (z,) = r.json()["shadowfax"]["ztp"]  # deduped by AWB
    assert z["penalty"] == 200.0
    r = _run(client, files, commit=True)
    assert r.status_code == 200, r.text
    n = db.execute(
        "SELECT COUNT(*) AS n FROM transactions WHERE event_type='ADJUSTMENT'"
    ).fetchone()["n"]
    assert n == 0  # not passed on unless chosen
    db.execute("DELETE FROM payout_order_days")
    db.execute("DELETE FROM company_cycles")
    db.commit()
    r = _run(client, files, commit=True, force="true", ztp_apply=json.dumps([ALPHA]))
    assert r.status_code == 200, r.text
    adj = db.execute(
        "SELECT amount, remarks FROM transactions WHERE event_type='ADJUSTMENT' AND person_id=?",
        (pids[ALPHA],),
    ).fetchone()
    assert adj["amount"] == -20_000 and "AWB1" in adj["remarks"]


def test_dues_owed_at_another_company_are_recovered_here(db, client):
    """The spec's Somnath: dues anchored at Kaptan, a Shadowfax payout now."""
    pids = _roster(db)
    alpha = pids[ALPHA]
    make_rider(db, alpha, "K-28158", "Kaptan", "Rider Alpha")
    db.execute(
        "UPDATE person_registry SET deduction_company='Kaptan', deduction_rider_id='K-28158' "
        "WHERE person_id=?",  # noqa: E501
        (alpha,),
    )
    db.execute(
        "INSERT INTO balances (person_id, current_balance) VALUES (?, ?) "
        "ON CONFLICT(person_id) DO UPDATE SET current_balance=excluded.current_balance",
        (alpha, -120_800),
    )
    db.commit()
    r = _run(client, fx.all_files())
    res = r.json()["result"]
    row = next(x for x in res["pay_rows"] + res["dues_rows"] if x["rider_id"] == ALPHA)
    assert row["payout"] == 1150.0
    assert row["released"] < 1150.0  # the ₹1,208 carried at Kaptan comes off this payout


def test_the_workbook_carries_the_detail_and_the_card(db, client):
    _roster(db)
    r = _run(client, fx.all_files(), commit=True)
    assert r.status_code == 200, r.text
    import base64

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(base64.b64decode(r.json()["xlsx"]["content_base64"])))
    assert "Shadowfax detail" in wb.sheetnames and "Ratecard used" in wb.sheetnames
    card = wb["Ratecard used"]
    pins = {
        str(row[0].value)
        for row in card.iter_rows(min_row=2)
        if row[0].value and row[0].value != "TOTAL"
    }
    assert "700154" in pins and "700033" in pins


def test_parse_sheet_suggests_the_cycle(db, client):
    r = client.post(
        "/api/cycles/parse-sheet",
        data={"company": "Shadowfax"},
        files=[("files", (n, b)) for n, b in fx.all_files()],
    )
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["kind"] == "shadowfax" and len(j["files"]) == 10
    assert (j["suggested_cycle_start"], j["suggested_cycle_end"]) == ("2026-09-17", "2026-09-27")


def test_the_ratecard_is_data(db, client):
    g = client.get("/api/companies/Shadowfax/ratecard").json()
    assert g["pincodes"] == 193
    one = next(x for x in g["rates"] if x["pincode"] == "700154")
    assert (one["ppd"], one["cod"], one["rvp"], one["sdd"], one["club"]) == (
        14.0,
        15.0,
        15.0,
        19.0,
        7.0,
    )
    import pandas as pd

    buf = io.BytesIO()
    pd.DataFrame(
        [
            {
                "cluster": "CCU_Dhakuria",
                "pincode": 700033,
                "RVP": 15,
                "COD": 15,
                "PPD": 14,
                "SDD": 19,
                "Club": 8,
            }
        ]
    ).to_excel(buf, index=False)
    pv = client.post(
        "/api/companies/Shadowfax/ratecard",
        data={"effective_from": "2026-09-26", "commit": "false"},
        files={"file": ("card.xlsx", buf.getvalue())},
    )
    assert pv.status_code == 200 and pv.json()["changed_count"] == 1 and not pv.json()["committed"]
    ok = client.post(
        "/api/companies/Shadowfax/ratecard",
        data={"effective_from": "2026-09-26", "commit": "true"},
        files={"file": ("card.xlsx", buf.getvalue())},
    )
    assert ok.status_code == 200 and ok.json()["committed"]
    # the new rate applies from its date only: 17 Sep keeps the old card, 26 Sep takes the new
    p = _price(db, fx.all_files())
    d = {(x["order_date"], x["pincode"]): x["rate_ppd"] for x in p.detail if x["rider_id"] == ALPHA}
    assert d[("2026-09-17", "700033")] == 1300 and d[("2026-09-26", "700033")] == 1400
    assert (
        db.execute(
            "SELECT COUNT(*) AS n FROM activity_log WHERE action='company.ratecard'"
        ).fetchone()["n"]
        == 1
    )
