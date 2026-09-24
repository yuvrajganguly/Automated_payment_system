"""Attaching a real company id to somebody we already have, by phone.

A Myntra run in September 2026 reported 17 onboardings and 16 riders falling
inactive. They were the same people. A recruiter had onboarded each of them in
the field before Myntra issued an id, so they sat on QSPEND placeholders; when
the payout file arrived with their real ids, nothing connected the two and the
operator was asked to onboard strangers who were already on the books.

A phone number is the only field in a payout file that survives a change of
rider id. These tests are mostly about when NOT to use it — linking the wrong
person attaches one rider's payout to another, which is worse than asking.
"""

from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook

from payout.domain.engine import process_cycle
from tests.conftest import make_person, make_rider

WEEK = date.today() - timedelta(days=date.today().weekday(), weeks=1)


def _file(rows, headers=("rider_id", "net_pay", "mobile")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _placeholder(db, name, phone, company="Kaptan", rider_id="QSPEND0031"):
    pid = make_person(db, name, balance=0, arrears=0)
    make_rider(db, pid, rider_id, company, name)
    db.execute("UPDATE rider_master SET mob_no=? WHERE rider_id=?", (phone, rider_id))
    db.commit()
    return pid


def _run(rows):
    return process_cycle("Kaptan", WEEK, WEEK + timedelta(days=6), _file(rows), commit=True)


# ── the case it exists for ───────────────────────────────────────────────────
def test_a_real_id_lands_on_the_person_who_was_waiting_for_it(db):
    pid = _placeholder(db, "Jit Dey", "7003664227")
    res = _run([("MNOW00471", 5000, "7003664227")])

    assert res.unknown_ids == [], "the rider was already on file"
    linked = [a for a in res.auto_linked if a["linked_from"] == "phone"]
    assert len(linked) == 1
    assert linked[0]["person_id"] == pid

    rows = {
        r["rider_id"]
        for r in db.execute(
            "SELECT rider_id FROM rider_master WHERE person_id=?", (pid,)
        ).fetchall()
    }
    assert rows == {"MNOW00471"}, "the placeholder should have been retired"


def test_the_payout_reaches_the_existing_person(db):
    """The point of all this: one person, one balance."""
    pid = _placeholder(db, "Jit Dey", "7003664227")
    _run([("MNOW00471", 5000, "7003664227")])
    assert (
        db.execute("SELECT COUNT(*) AS n FROM transactions WHERE person_id=?", (pid,)).fetchone()[
            "n"
        ]
        > 0
    )
    assert db.execute("SELECT COUNT(*) AS n FROM person_registry").fetchone()["n"] == 1, (
        "a second person was created anyway"
    )


def test_a_phone_written_with_a_country_code_still_matches(db):
    _placeholder(db, "Jit Dey", "+91 70036-64227")
    res = _run([("MNOW00471", 5000, "07003664227")])
    assert res.unknown_ids == []


# ── when it must not fire ────────────────────────────────────────────────────
def test_two_people_on_one_handset_are_left_for_a_human(db):
    """Brothers sharing a phone is ordinary. Guessing here hands one man's
    payout to the other."""
    _placeholder(db, "Jit Dey", "7003664227", rider_id="QSPEND0031")
    _placeholder(db, "Raju Dey", "7003664227", rider_id="QSPEND0032")
    res = _run([("MNOW00471", 5000, "7003664227")])
    assert res.unknown_ids == ["MNOW00471"]


def test_a_name_that_is_plainly_somebody_else_blocks_the_link(db):
    """A phone typed into the wrong row should not move money."""
    _placeholder(db, "Jit Dey", "7003664227")
    res = _run(
        [("MNOW00471", 5000, "7003664227")],
    )
    assert res.unknown_ids == []  # control: no name column, so no contradiction

    _placeholder(db, "Somnath Sardar", "9000000001", rider_id="QSPEND0040")
    wb = Workbook()
    ws = wb.active
    ws.append(["rider_id", "net_pay", "mobile", "rider_name"])
    ws.append(["MNOW00999", 5000, "9000000001", "Milon Ghosh"])
    buf = io.BytesIO()
    wb.save(buf)
    res = process_cycle(
        "Kaptan", WEEK + timedelta(days=7), WEEK + timedelta(days=13), buf.getvalue(), commit=True
    )
    assert res.unknown_ids == ["MNOW00999"]


def test_somebody_who_already_has_a_real_id_here_is_not_given_a_second(db):
    """They may legitimately hold two, but that is a decision somebody makes,
    not one inferred from a phone number."""
    pid = make_person(db, "Already Here", balance=0, arrears=0)
    make_rider(db, pid, "MNOW00001", "Kaptan", "Already Here")
    db.execute("UPDATE rider_master SET mob_no='7003664227' WHERE rider_id='MNOW00001'")
    db.commit()
    res = _run([("MNOW00471", 5000, "7003664227")])
    assert res.unknown_ids == ["MNOW00471"]


def test_no_phone_in_the_file_changes_nothing(db):
    """Most companies do not send one. The old behaviour has to survive."""
    _placeholder(db, "Jit Dey", "7003664227")
    wb = Workbook()
    ws = wb.active
    ws.append(["rider_id", "net_pay"])
    ws.append(["MNOW00471", 5000])
    buf = io.BytesIO()
    wb.save(buf)
    res = process_cycle("Kaptan", WEEK, WEEK + timedelta(days=6), buf.getvalue(), commit=True)
    assert res.unknown_ids == ["MNOW00471"]


def test_an_unrecognised_phone_is_still_unknown(db):
    _placeholder(db, "Jit Dey", "7003664227")
    res = _run([("MNOW00471", 5000, "9999999999")])
    assert res.unknown_ids == ["MNOW00471"]


def test_the_link_is_reported_so_the_preview_can_show_it(db):
    """Nothing here is silent — this proposes, the operator commits."""
    _placeholder(db, "Jit Dey", "7003664227")
    res = process_cycle(
        "Kaptan",
        WEEK,
        WEEK + timedelta(days=6),
        _file([("MNOW00471", 5000, "7003664227")]),
        commit=False,
    )
    assert any("matched Jit Dey by phone" in w for w in res.warnings)
    assert res.auto_linked[0]["retired_placeholders"] == ["QSPEND0031"]


def test_a_refused_match_still_tells_the_operator_who_it_might_be(db):
    """Refusing to link is right when it is not certain. Leaving a bare id and
    a payout on screen is how seventeen "new" riders got created for sixteen
    people we already had — so the modal gets the candidates."""
    _placeholder(db, "Jit Dey", "7003664227", rider_id="QSPEND0031")
    _placeholder(db, "Raju Dey", "7003664227", rider_id="QSPEND0032")
    res = _run([("MNOW00471", 5000, "7003664227")])

    assert res.unknown_ids == ["MNOW00471"], "two on one handset: still refused"
    row = res.unknown_riders[0]
    names = {m["display_name"] for m in row["possible_matches"]}
    assert names == {"Jit Dey", "Raju Dey"}
    assert all("phone=7003664227" in m["evidence"] for m in row["possible_matches"])


def test_a_genuinely_new_rider_gets_no_suggestions(db):
    _placeholder(db, "Jit Dey", "7003664227")
    res = _run([("MNOW00999", 5000, "8888888888")])
    assert res.unknown_riders[0]["possible_matches"] == []
